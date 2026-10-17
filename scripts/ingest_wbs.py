#!/usr/bin/env python3
"""Ingest the Hamtayar WBS workbook into a normalized, reproducible export.

Usage:
    python scripts/ingest_wbs.py --source PATH --out DIR

Reads the original workbook (never modifies it), validates structure/counts,
and writes JSON + CSV files under DIR. Exits non-zero on validation failure.
"""
import argparse
import csv
import json
import os
import sys

EXPECTED_SHEETS = ["Backlog (WBS)", "خلاصه اپیک‌ها", "سؤالات باز از کارفرما", "راهنمای رنگ‌ها"]
EXPECTED = {
    "epics": 24,
    "features": 98,
    "tasks": 442,
    "task_phase": {"فاز ۰": 62, "فاز ۱": 276, "فاز ۲": 51, "فاز ۳": 29,
                   "فاز ۴": 9, "فاز ۵": 5, "پیشنهادی": 10},
    "task_prio": {"P0": 313, "P1": 90, "P2": 39},
}


def load_workbook(path):
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
        sys.exit(2)
    if not os.path.isfile(path):
        print(f"ERROR: source workbook not found: {path}", file=sys.stderr)
        sys.exit(2)
    return openpyxl.load_workbook(path, data_only=True)


def ingest(source):
    wb = load_workbook(source)
    # sheet presence
    for s in EXPECTED_SHEETS:
        if s not in wb.sheetnames:
            print(f"ERROR: missing sheet {s!r}. Found: {wb.sheetnames}", file=sys.stderr)
            sys.exit(1)

    wbs = wb["Backlog (WBS)"]
    rows = list(wbs.iter_rows(min_row=2, values_only=True))
    header = [c for c in (r[0] for r in [])]  # placeholder
    # actual header
    hdr = [c.value for c in wbs[1]]
    records = []
    for r in rows:
        if all(v is None for v in r):
            continue
        rec = {
            "wbs_id": r[0],
            "type": r[1],
            "title": r[2],
            "phase": r[3],
            "priority": r[4],
            "owner": r[5],
            "status": r[6],
            "estimate_days": r[7],
            "dependencies": r[8],
            "notes_acceptance": r[9],
            "waiting_decision": r[10],
        }
        records.append(rec)

    epics = [r for r in records if r["type"] == "Epic"]
    features = [r for r in records if r["type"] == "Feature"]
    tasks = [r for r in records if r["type"] == "Task"]

    # validation
    errors = []
    if len(epics) != EXPECTED["epics"]:
        errors.append(f"epic count {len(epics)} != {EXPECTED['epics']}")
    if len(features) != EXPECTED["features"]:
        errors.append(f"feature count {len(features)} != {EXPECTED['features']}")
    if len(tasks) != EXPECTED["tasks"]:
        errors.append(f"task count {len(tasks)} != {EXPECTED['tasks']}")

    from collections import Counter
    phase_c = Counter(t["phase"] for t in tasks)
    for k, v in EXPECTED["task_phase"].items():
        if phase_c.get(k, 0) != v:
            errors.append(f"task phase {k!r}: {phase_c.get(k,0)} != {v}")
    prio_c = Counter(t["priority"] for t in tasks)
    for k, v in EXPECTED["task_prio"].items():
        if prio_c.get(k, 0) != v:
            errors.append(f"task priority {k!r}: {prio_c.get(k,0)} != {v}")

    if errors:
        print("VALIDATION FAILED:", file=sys.stderr)
        for e in errors:
            print("  - " + e, file=sys.stderr)
        sys.exit(1)

    # epic summary sheet
    es = wb["خلاصه اپیک‌ها"]
    epic_summary = []
    for r in es.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in r):
            continue
        epic_summary.append({"wbs": r[0], "epic": r[1], "features": r[2], "tasks": r[3]})

    # open questions sheet
    oq = wb["سؤالات باز از کارفرما"]
    questions = []
    for r in oq.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in r):
            continue
        questions.append({"id": r[0], "topic": r[1], "question": r[2], "impact": r[3]})

    # color legend
    cl = wb["راهنمای رنگ‌ها"]
    legend = []
    for r in cl.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in r):
            continue
        legend.append({"status": r[0], "meaning": r[1]})

    return {
        "header": hdr,
        "records": records,
        "epics": epics,
        "features": features,
        "tasks": tasks,
        "epic_summary": epic_summary,
        "open_questions": questions,
        "color_legend": legend,
        "counts": {
            "epics": len(epics),
            "features": len(features),
            "tasks": len(tasks),
            "task_phase": dict(phase_c),
            "task_priority": dict(prio_c),
        },
    }


def write_outputs(data, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    # JSON
    with open(os.path.join(out_dir, "wbs.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    # tasks.csv
    task_fields = ["wbs_id", "type", "title", "phase", "priority", "owner", "status",
                   "estimate_days", "dependencies", "notes_acceptance", "waiting_decision"]
    with open(os.path.join(out_dir, "tasks.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=task_fields)
        w.writeheader()
        for t in data["tasks"]:
            w.writerow({k: t.get(k) for k in task_fields})
    # features.csv
    with open(os.path.join(out_dir, "features.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=task_fields)
        w.writeheader()
        for t in data["features"]:
            w.writerow({k: t.get(k) for k in task_fields})
    # epics.csv
    with open(os.path.join(out_dir, "epics.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=task_fields)
        w.writeheader()
        for t in data["epics"]:
            w.writerow({k: t.get(k) for k in task_fields})
    # open questions json
    with open(os.path.join(out_dir, "open_questions.json"), "w", encoding="utf-8") as f:
        json.dump(data["open_questions"], f, ensure_ascii=False, indent=2)
    # counts
    with open(os.path.join(out_dir, "counts.json"), "w", encoding="utf-8") as f:
        json.dump(data["counts"], f, ensure_ascii=False, indent=2)
    print(f"Wrote normalized export to {out_dir}")
    print(json.dumps(data["counts"], ensure_ascii=True, indent=2))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True, help="Path to Hamtayar_Backlog_WBS.xlsx")
    ap.add_argument("--out", required=True, help="Output directory")
    args = ap.parse_args()
    data = ingest(args.source)
    write_outputs(data, args.out)


if __name__ == "__main__":
    main()
