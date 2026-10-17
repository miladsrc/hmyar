"""Tests for WBS ingestion."""
import json
import os
import subprocess
import sys
import tempfile

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import ingest_wbs

REAL_WBS = r"C:\Users\m.barani\Downloads\Hamtayar_Backlog_WBS.xlsx"


def make_fixture(path, epics=2, features=3, tasks=5):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Backlog (WBS)"
    ws.append(["WBS ID", "نوع", "عنوان (ساختار درختی)", "فاز", "اولویت پیشنهادی",
               "مسئول", "وضعیت", "تخمین (نفر-روز)", "وابستگی", "یادداشت / معیار پذیرش", "منتظر چه تصمیمی؟"])
    # epics
    for i in range(1, epics + 1):
        ws.append([str(i), "Epic", f"Epic {i}", None, None, None, None, None, None, None, None])
    # features
    fid = 0
    for e in range(1, epics + 1):
        for _ in range(features // epics or 1):
            fid += 1
            ws.append([f"{e}.{fid}", "Feature", f"Feature {fid}", "فاز ۰", "P0", None, None, None, None, None, None])
    # tasks
    tid = 0
    for e in range(1, epics + 1):
        for _ in range(tasks // epics or 1):
            tid += 1
            ws.append([f"{e}.{tid}.{tid}", "Task", f"Task {tid}", "فاز ۰", "P0", None, "انجام نشده", None, None, None, None])
    # other sheets
    wb.create_sheet("خلاصه اپیک‌ها").append(["WBS", "اپیک", "تعداد Feature", "تعداد Task"])
    wb.create_sheet("سؤالات باز از کارفرما").append(["#", "موضوع", "سؤال / ابهام", "تأثیر روی بک‌لاگ"])
    wb.create_sheet("راهنمای رنگ‌ها").append(["وضعیت", "معنی"])
    wb.save(path)


def test_ingest_real_wbs_counts():
    if not os.path.isfile(REAL_WBS):
        pytest.skip("real WBS not present")
    data = ingest_wbs.ingest(REAL_WBS)
    assert data["counts"]["epics"] == 24
    assert data["counts"]["features"] == 98
    assert data["counts"]["tasks"] == 442
    assert data["counts"]["task_phase"]["فاز ۰"] == 62
    assert data["counts"]["task_phase"]["پیشنهادی"] == 10
    assert data["counts"]["task_priority"]["P0"] == 313


def test_ingest_missing_source_exits():
    with pytest.raises(SystemExit):
        ingest_wbs.ingest("C:/nope/missing.xlsx")


def test_ingest_bad_counts_fail_validation():
    # ingest() validates against the Hamtayar expected baseline; a mismatched workbook
    # must fail loudly (SystemExit) rather than silently producing wrong data.
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "t.xlsx")
        make_fixture(p, epics=2, features=4, tasks=6)
        with pytest.raises(SystemExit):
            ingest_wbs.ingest(p)


def test_cli_writes_outputs():
    if not os.path.isfile(REAL_WBS):
        pytest.skip("real WBS not present")
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "out")
        subprocess.run([sys.executable, os.path.join(ROOT, "scripts", "ingest_wbs.py"),
                        "--source", REAL_WBS, "--out", out], check=True)
        assert os.path.isfile(os.path.join(out, "wbs.json"))
        assert os.path.isfile(os.path.join(out, "tasks.csv"))
        with open(os.path.join(out, "counts.json"), encoding="utf-8") as f:
            c = json.load(f)
        assert c["tasks"] == 442